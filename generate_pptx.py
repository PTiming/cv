# USAGE
# python generate_pptx.py --output presentation.pptx --title "Face Liveness Detection"

# import the necessary packages
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import argparse

def add_title_slide(prs, title, subtitle):
    """Add a title slide to the presentation."""
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)
    
    # Add title
    left = Inches(0.5)
    top = Inches(2.5)
    width = Inches(9)
    height = Inches(1.5)
    
    title_box = slide.shapes.add_textbox(left, top, width, height)
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = title
    title_para.font.size = Pt(44)
    title_para.font.bold = True
    title_para.alignment = PP_ALIGN.CENTER
    
    # Add subtitle
    sub_top = Inches(4)
    sub_box = slide.shapes.add_textbox(left, sub_top, width, Inches(1))
    sub_frame = sub_box.text_frame
    sub_para = sub_frame.paragraphs[0]
    sub_para.text = subtitle
    sub_para.font.size = Pt(24)
    sub_para.alignment = PP_ALIGN.CENTER
    
    return slide


def add_content_slide(prs, title, bullet_points):
    """Add a content slide with title and bullet points."""
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)
    
    # Add title
    left = Inches(0.5)
    top = Inches(0.5)
    width = Inches(9)
    height = Inches(1)
    
    title_box = slide.shapes.add_textbox(left, top, width, height)
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = title
    title_para.font.size = Pt(36)
    title_para.font.bold = True
    
    # Add bullet points
    content_top = Inches(1.75)
    content_height = Inches(5)
    
    content_box = slide.shapes.add_textbox(left, content_top, width, content_height)
    content_frame = content_box.text_frame
    content_frame.word_wrap = True
    
    for i, point in enumerate(bullet_points):
        if i == 0:
            para = content_frame.paragraphs[0]
        else:
            para = content_frame.add_paragraph()
        
        para.text = f"• {point}"
        para.font.size = Pt(24)
        para.space_after = Pt(12)
    
    return slide


def add_code_slide(prs, title, code_text):
    """Add a slide with code snippet."""
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)
    
    # Add title
    left = Inches(0.5)
    top = Inches(0.3)
    width = Inches(9)
    height = Inches(0.8)
    
    title_box = slide.shapes.add_textbox(left, top, width, height)
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = title
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    
    # Add code box
    code_top = Inches(1.3)
    code_height = Inches(5.5)
    
    code_box = slide.shapes.add_textbox(left, code_top, width, code_height)
    code_frame = code_box.text_frame
    code_frame.word_wrap = True
    
    code_para = code_frame.paragraphs[0]
    code_para.text = code_text
    code_para.font.size = Pt(14)
    code_para.font.name = "Courier New"
    
    return slide


def create_liveness_presentation(output_path, title):
    """Create a presentation about the Face Liveness Detection project."""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Slide 1: Title Slide
    add_title_slide(
        prs,
        title,
        "Detecting Real vs Fake Faces Using Deep Learning"
    )
    
    # Slide 2: Project Overview
    add_content_slide(
        prs,
        "Project Overview",
        [
            "Face liveness detection using deep learning",
            "Distinguishes between real faces and spoofing attempts",
            "Uses OpenCV for face detection",
            "Keras/TensorFlow for liveness classification",
            "Supports video files and real-time webcam input"
        ]
    )
    
    # Slide 3: Problem Statement
    add_content_slide(
        prs,
        "Problem Statement",
        [
            "Face recognition systems are vulnerable to spoofing attacks",
            "Attackers can use photos, videos, or masks to bypass security",
            "Liveness detection adds an extra layer of security",
            "Determines if the face is from a live person",
            "Essential for secure authentication systems"
        ]
    )
    
    # Slide 4: Technical Approach
    add_content_slide(
        prs,
        "Technical Approach",
        [
            "Step 1: Gather training data (real and fake face images)",
            "Step 2: Train a CNN model to classify real vs fake",
            "Step 3: Use OpenCV DNN for face detection",
            "Step 4: Apply liveness model to detected faces",
            "Step 5: Display results with confidence scores"
        ]
    )
    
    # Slide 5: Dataset Collection
    add_content_slide(
        prs,
        "Dataset Collection",
        [
            "Real faces: Videos of live people",
            "Fake faces: Videos of photos/screens showing faces",
            "Face detector extracts face regions from video frames",
            "Images are cropped and saved to dataset folders",
            "Balanced dataset with equal real and fake samples"
        ]
    )
    
    # Slide 6: Model Architecture
    add_content_slide(
        prs,
        "Model Architecture",
        [
            "Convolutional Neural Network (CNN)",
            "Input: 32x32 RGB face images",
            "Multiple Conv2D layers with ReLU activation",
            "MaxPooling for spatial reduction",
            "Dense layers for classification",
            "Binary output: Real (1) or Fake (0)"
        ]
    )
    
    # Slide 7: Key Technologies
    add_content_slide(
        prs,
        "Key Technologies",
        [
            "Python 3.8+ - Programming language",
            "OpenCV - Computer vision and face detection",
            "Keras/TensorFlow - Deep learning framework",
            "NumPy - Numerical computations",
            "imutils - Image processing utilities"
        ]
    )
    
    # Slide 8: Pipeline Workflow
    add_content_slide(
        prs,
        "Pipeline Workflow",
        [
            "1. Install dependencies (01_install.bat)",
            "2. Gather face examples (02_gather.bat)",
            "3. Train liveness model (03_trainLiveness.bat)",
            "4. Run on video files (04_runLiveness.bat)",
            "5. Run with webcam (05_webcam.bat)"
        ]
    )
    
    # Slide 9: Code Example - Face Detection
    code_example = """# Load face detector
protoPath = os.path.sep.join([detector, "deploy.prototxt"])
modelPath = os.path.sep.join([detector, 
    "res10_300x300_ssd_iter_140000.caffemodel"])
net = cv2.dnn.readNetFromCaffe(protoPath, modelPath)

# Detect faces in frame
blob = cv2.dnn.blobFromImage(frame, 1.0, (300, 300),
    (104.0, 177.0, 123.0))
net.setInput(blob)
detections = net.forward()"""
    
    add_code_slide(prs, "Code: Face Detection", code_example)
    
    # Slide 10: Code Example - Liveness Prediction
    liveness_code = """# Preprocess face for liveness model
face = cv2.resize(face, (32, 32))
face = face.astype("float") / 255.0
face = img_to_array(face)
face = np.expand_dims(face, axis=0)

# Predict liveness
preds = model.predict(face)[0]
label = le.classes_[np.argmax(preds)]
confidence = preds[np.argmax(preds)]"""
    
    add_code_slide(prs, "Code: Liveness Prediction", liveness_code)
    
    # Slide 11: Results Visualization
    add_content_slide(
        prs,
        "Results Visualization",
        [
            "Real faces: Green bounding box",
            "Fake faces: Red bounding box",
            "Confidence score displayed above face",
            "Real-time processing on video stream",
            "Output saved as annotated video file"
        ]
    )
    
    # Slide 12: Applications
    add_content_slide(
        prs,
        "Applications",
        [
            "Mobile banking and payment authentication",
            "Access control systems",
            "Online identity verification",
            "Attendance management systems",
            "Social media account security"
        ]
    )
    
    # Slide 13: Future Improvements
    add_content_slide(
        prs,
        "Future Improvements",
        [
            "Multi-frame temporal analysis",
            "3D depth sensing integration",
            "Challenge-response mechanisms (blink, smile)",
            "Edge device optimization",
            "Transfer learning with larger datasets"
        ]
    )
    
    # Slide 14: Conclusion
    add_content_slide(
        prs,
        "Conclusion",
        [
            "Liveness detection enhances face recognition security",
            "Deep learning provides accurate real vs fake classification",
            "Simple yet effective CNN architecture",
            "Easy to integrate with existing systems",
            "Open for further improvements and extensions"
        ]
    )
    
    # Slide 15: Thank You
    add_title_slide(
        prs,
        "Thank You!",
        "Questions?"
    )
    
    # Save the presentation
    try:
        prs.save(output_path)
        print(f"[INFO] Presentation saved to: {output_path}")
    except PermissionError:
        print(f"[ERROR] Permission denied when saving to: {output_path}")
        return None
    except OSError as e:
        print(f"[ERROR] Failed to save presentation: {e}")
        return None
    
    return output_path


def main():
    # Construct the argument parser and parse the arguments
    ap = argparse.ArgumentParser()
    ap.add_argument("-o", "--output", type=str, default="liveness_presentation.pptx",
        help="path to output PowerPoint file")
    ap.add_argument("-t", "--title", type=str, default="Face Liveness Detection",
        help="title of the presentation")
    args = vars(ap.parse_args())
    
    # Create the presentation
    create_liveness_presentation(args["output"], args["title"])


if __name__ == "__main__":
    main()
